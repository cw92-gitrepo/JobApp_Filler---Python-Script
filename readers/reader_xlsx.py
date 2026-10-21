"""Reader for .xlsx files: raw extraction only.

"Excel file in, plain tables out." Validation and building the Resume
happen in resume_reader.py.
"""

from pathlib import Path
from zipfile import BadZipFile

import openpyxl
from openpyxl.utils.exceptions import InvalidFileException

from readers.resume_reader import Row, Table, Tables
from util.errors import ResumeDataError

# One tab per section, returned in the contract's fixed order.
SHEET_ORDER = ("General", "Personal", "Employment", "Education")


def read(path: Path) -> Tables:
    # read_only: streams rows instead of loading the whole file into memory.
    # data_only: a formula cell gives its last computed value, not "=A1+B1".
    try:
        workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except PermissionError:
        # The usual cause on Windows: the file is open in Excel.
        raise ResumeDataError(f"Can't open {path.name}. Close it in Excel and try again.") from None
    except (InvalidFileException, BadZipFile, OSError, KeyError) as error:
        raise ResumeDataError(f"{path.name} isn't a readable .xlsx file ({error}).") from None

    try:
        missing = [name for name in SHEET_ORDER if name not in workbook.sheetnames]
        if missing:
            raise ResumeDataError(f"{path.name} is missing the tab(s): {', '.join(missing)}.")
        general, personal, employment, education = (_sheet_to_table(workbook[name]) for name in SHEET_ORDER)
    finally:
        workbook.close()  # read-only workbooks keep the file open until closed
    return general, personal, employment, education


def _sheet_to_table(sheet) -> Table:
    rows: list[Row] = [tuple(row) for row in sheet.iter_rows(values_only=True)]
    # Excel quirk: formatted-but-empty rows come back as all-None. Drop them.
    return tuple(row for row in rows if any(cell is not None for cell in row))
