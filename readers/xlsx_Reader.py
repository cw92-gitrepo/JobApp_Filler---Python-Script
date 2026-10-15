from openpyxl import load_workbook
from pathlib import Path
from base_reader import Reader
from data.resume import Resume

class Xlsx_Reader(Reader):

    def read_from_file(self, path:Path):
        workbook = load_workbook(path, read_only=True, data_only=True)

        for sheet in workbook.worksheets:
            match sheet.title:
                case "Personal":
                    fill_resume = self.make_usr_personal
                case "General":
                    fill_resume = self.make_usr_general
                case "Employment":
                    fill_resume = self.make_employment
                case "Education":
                    fill_resume = self.make_education


            for row in sheet.iter_rows(values_only=True):
                fill_resume(row)


        return Resume(self.usr_personal, self.usr_general, self.usr_employment, self.usr_education)

        pass

    
    


    pass